"""Template-based DOCX drafts and text-based PDF exports for simple resumes."""

from collections import Counter
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import re
from urllib.parse import urlsplit
from xml.sax.saxutils import escape, quoteattr

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.hyperlink import Hyperlink
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, HRFlowable, PageTemplate, Paragraph


EXPERIENCE_HEADING = re.compile(r"\b(experience|employment|work history|career history)\b", re.IGNORECASE)


def safe_link(url):
    try:
        return urlsplit(url).scheme.lower() in {"https", "http", "mailto", "tel"}
    except ValueError:
        return False


def original_links(document):
    return [link for paragraph in document.paragraphs for link in paragraph.hyperlinks
            if link.text and safe_link(link.url)]


def normalize_resume_punctuation(document):
    """Use plain punctuation in visible text without touching formatting or URLs."""
    replacements = str.maketrans({"—": "-", "–": "-", "‘": "'", "’": "'",
                                 "“": '"', "”": '"', "…": "..."})
    for node in document.element.body.xpath(".//w:t"):
        if node.text:
            node.text = node.text.translate(replacements)


def save_resume(draft, path, original_bytes):
    """Keep the original top block and template styles; rewrite only the body."""
    document = Document(BytesIO(original_bytes))
    paragraphs = document.paragraphs
    headings = [p for p in paragraphs if p.style.name.lower().startswith("heading")]
    first_heading = paragraphs.index(headings[0]) if headings else 0
    header = paragraphs[:first_heading] if 0 < first_heading <= 10 else []
    heading_template = headings[0] if headings else None
    body_template = next((p for p in paragraphs[first_heading + 1:]
                          if p.text.strip() and p not in headings), None)
    bullet_template = next((p for p in paragraphs if p._p.pPr is not None
                            and p._p.pPr.numPr is not None), None)
    links = original_links(document)
    original_paragraphs = {p.text: p for p in paragraphs if p.text.strip()}
    counts = Counter(link.text for link in links)
    unique_links = {link.text: link for link in links if counts[link.text] == 1}
    used = {link.url for paragraph in header for link in paragraph.hyperlinks}
    pattern = re.compile("|".join(re.escape(label) for label in sorted(unique_links, key=len, reverse=True))) if unique_links else None

    body = document.element.body
    for element in list(body):
        if element.tag != qn("w:sectPr") and all(element is not p._p for p in header):
            body.remove(element)

    def add(text, template=None, style=None, *, bold=False):
        paragraph = document.add_paragraph(style=style)
        if text in original_paragraphs:
            original = original_paragraphs[text]
            copied = deepcopy(original._p)
            if bold:
                for run in copied.xpath(".//w:r"):
                    run.get_or_add_rPr().get_or_add_b().val = True
            paragraph._p.addprevious(copied)
            body.remove(paragraph._p)
            used.update(link.url for link in original.hyperlinks)
            return
        if template is not None and template._p.pPr is not None:
            if paragraph._p.pPr is not None:
                paragraph._p.remove(paragraph._p.pPr)
            paragraph._p.insert(0, deepcopy(template._p.pPr))
        run_template = next((r for r in template.runs if r.text), None) if template is not None else None

        def add_text(value):
            run = paragraph.add_run(value)
            if run_template is not None and run_template._r.rPr is not None:
                run._r.insert(0, deepcopy(run_template._r.rPr))

        offset = 0
        for match in pattern.finditer(text) if pattern else []:
            add_text(text[offset:match.start()])
            link = unique_links[match.group()]
            paragraph._p.append(deepcopy(link._hyperlink))
            used.add(link.url)
            offset = match.end()
        add_text(text[offset:])
        if bold:
            for run in paragraph._p.xpath(".//w:r"):
                run.get_or_add_rPr().get_or_add_b().val = True
        return paragraph

    if not header:
        add(draft.name, style="Title")
        for line in draft.contact_lines:
            add(line, body_template)
    add("Professional Summary", heading_template, "Heading 1")
    add(draft.summary, body_template)
    for section in draft.sections:
        add(section.heading, heading_template, "Heading 1")
        # The draft contract uses plain lines for role metadata and '- ' for achievements.
        is_experience = bool(EXPERIENCE_HEADING.search(section.heading))
        for text in section.paragraphs:
            if text.startswith("- "):
                add(text[2:], bullet_template, "List Bullet" if bullet_template is None else None)
            else:
                add(text, body_template, bold=is_experience)
    # Never silently lose original project/profile links when the LLM omits their labels.
    remaining = [link for link in links if link.url not in used]
    if remaining:
        add("Links", heading_template, "Heading 1")
        for link in remaining:
            if link.url not in used:
                paragraph = document.add_paragraph()
                paragraph._p.append(deepcopy(link._hyperlink))
                used.add(link.url)
    normalize_resume_punctuation(document)
    document.save(path)


def pdf_fonts():
    """Prefer local Calibri, then Arial; bundled Vera is the portable fallback."""
    roots = [Path.home() / "Library/Fonts", Path("/Library/Fonts"),
             Path("/System/Library/Fonts/Supplemental"), Path("C:/Windows/Fonts")]
    families = [("Calibri", ("calibri.ttf", "calibrib.ttf", "calibrii.ttf", "calibriz.ttf")),
                ("Arial", ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf")),
                ("Arial", ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"))]
    for family, filenames in families:
        for root in roots:
            paths = [root / filename for filename in filenames]
            if all(path.is_file() for path in paths):
                break
        else:
            continue
        break
    else:
        family = "Vera"
        root = next(Path(p) for p in rl_config.TTFSearchPath if (Path(p) / "Vera.ttf").is_file())
        paths = [root / name for name in ("Vera.ttf", "VeraBd.ttf", "VeraIt.ttf", "VeraBI.ttf")]
    names = [f"Jobbing-{family}-{suffix}" for suffix in ("regular", "bold", "italic", "bolditalic")]
    for name, path in zip(names, paths):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    return names


def font_value(run, paragraph, attribute, default):
    fonts = [run.font]
    if run.style is not None:
        fonts.append(run.style.font)
    fonts.append(paragraph.style.font)
    style = paragraph.style.base_style
    while style is not None:
        fonts.append(style.font)
        style = style.base_style
    for font in fonts:
        value = font.color.rgb if attribute == "color" else getattr(font, attribute)
        if value is not None:
            return value
    return default


def save_pdf(docx_path, pdf_path, *, resume=True):
    """Typeset a readable resume without changing the editable DOCX.

    Set resume=False to retain paragraph styling for motivation letters. This is
    not a Word layout engine: floating objects, columns and headers are omitted.
    """
    document = Document(docx_path)
    if resume:
        normalize_resume_punctuation(document)
    fonts = pdf_fonts()
    section = document.sections[0]
    width, height = section.page_width.pt, section.page_height.pt
    left, right = section.left_margin.pt, section.right_margin.pt
    top, bottom = section.top_margin.pt, section.bottom_margin.pt
    if resume:
        left = right = 50.4  # 0.7-inch margins, with room for readable line lengths.
        top, bottom = 43.2, 46.8
    pdf = BaseDocTemplate(str(pdf_path), pagesize=(width, height), title=Path(docx_path).stem,
                          author="", leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom)
    frame = Frame(left, bottom, width - left - right, height - top - bottom,
                  leftPadding=0, rightPadding=0, topPadding=0, bottomPadding=0)

    def footer(canvas, doc):
        if not resume:
            return
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D8E2DC"))
        canvas.setLineWidth(0.5)
        canvas.line(left, bottom - 13, width - right, bottom - 13)
        canvas.setFont(fonts[0], 8)
        canvas.setFillColor(colors.HexColor("#52616B"))
        canvas.drawRightString(width - right, bottom - 27, str(doc.page))
        canvas.restoreState()

    pdf.addPageTemplates(PageTemplate(id="resume", frames=[frame], onPageEnd=footer))
    # Size, leading, color, space before, space after. No source-template shrinking.
    typography = {
        "name": (28, 32, "172B35", 0, 5),
        "subtitle": (12, 16, "16634B", 0, 5),
        "contact": (9.5, 13, "52616B", 0, 2),
        "heading": (11, 14, "16634B", 14, 4),
        "role": (11, 14.5, "172B35", 8, 4),
        "body": (10.5, 14.25, "293640", 0, 5),
        "bullet": (10.5, 14.25, "293640", 0, 3.5),
    }
    paragraphs = [p for p in document.paragraphs if p.text.strip() or not resume]
    header_end = next((i for i, p in enumerate(paragraphs)
                       if p.style.name.lower().startswith("heading")), 0)
    story = []
    if resume and header_end:
        accent = HRFlowable(width=42, thickness=3, color=colors.HexColor("#1E8E3E"),
                            hAlign="LEFT", spaceAfter=10)
        accent.keepWithNext = True
        story.append(accent)
    in_experience, in_skills, previous_kind = False, False, None
    for index, paragraph in enumerate(paragraphs):
        sizes = []
        is_heading = paragraph.style.name.lower().startswith("heading")
        bullet = paragraph.style.name == "List Bullet" or (paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None)
        if is_heading:
            in_experience = bool(EXPERIENCE_HEADING.search(paragraph.text))
            in_skills = bool(re.search(r"\bskills\b", paragraph.text, re.IGNORECASE))
        kind = "heading" if is_heading else "bullet" if bullet else "role" if in_experience else "body"
        if index < header_end:
            # Only treat an existing larger header line as a subtitle, never invent one.
            large_header = any(font_value(r, paragraph, "size", Pt(0)) >= Pt(12)
                               for r in paragraph.runs if r.text.strip())
            kind = "name" if index == 0 else "subtitle" if large_header else "contact"
        size, leading, color, before, after = typography[kind]
        label_length = paragraph.text.find(":") + 1 if resume and in_skills and kind == "body" else 0
        label_remaining = label_length if label_length <= 40 else 0

        def markup(run, *, link=False):
            nonlocal label_remaining
            run_size = font_value(run, paragraph, "size", None)
            run_size = size if resume else run_size.pt if run_size is not None else 10.5
            sizes.append(run_size)
            bold = bool(font_value(run, paragraph, "bold", False))
            italic = bool(font_value(run, paragraph, "italic", False))
            if resume and kind in {"name", "heading", "role", "subtitle", "contact"}:
                bold = kind in {"name", "heading", "role"}
                italic = False
            font = fonts[int(bold) + 2 * int(italic)]
            run_color = ("16634B" if link else color) if resume else str(font_value(run, paragraph, "color", "000000"))
            emphasized = min(label_remaining, len(run.text))
            label_remaining -= emphasized
            text = escape(run.text)
            if emphasized:
                text = f'<font name="{fonts[1]}">{escape(run.text[:emphasized])}</font>{escape(run.text[emphasized:])}'
            text = text.replace("\n", "<br/>").replace("\t", "    ")
            underline = link if resume else font_value(run, paragraph, "underline", False)
            if underline:
                text = f"<u>{text}</u>"
            return f'<font name="{font}" size="{run_size}" color="#{run_color}">{text}</font>'

        fragments = []
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink):
                text = "".join(markup(run, link=safe_link(item.url)) for run in item.runs)
                fragments.append(f'<link href={quoteattr(item.url)}>{text}</link>' if safe_link(item.url) else text)
            else:
                fragments.append(markup(item))
        text = "".join(fragments) or "<br/>"
        if resume:
            if kind == "role" and previous_kind in {"heading", "role"}:
                before = 0
            style = ParagraphStyle(
                "resume", fontName=fonts[0], fontSize=size, leading=leading,
                textColor=colors.HexColor(f"#{color}"), spaceBefore=before, spaceAfter=after,
                leftIndent=13 if bullet else 0, bulletIndent=1, bulletFontName=fonts[0],
                bulletFontSize=8, bulletColor=colors.HexColor("#16634B"),
                keepWithNext=kind in {"name", "subtitle", "contact", "heading", "role"},
                allowWidows=0, allowOrphans=0,
            )
        else:
            size = max(sizes, default=10.5)
            fmt = paragraph.paragraph_format
            base = paragraph.style.paragraph_format

            def setting(name, default=None):
                value = getattr(fmt, name)
                return value if value is not None else getattr(base, name) if getattr(base, name) is not None else default

            line = setting("line_spacing", 1.15)
            leading = line.pt if hasattr(line, "pt") else float(line) * size
            indent = setting("left_indent")
            indent = indent.pt if indent is not None else 14 if bullet else 0
            before, after = setting("space_before"), setting("space_after")
            style = ParagraphStyle(
                "document", fontName=fonts[0], fontSize=size, leading=max(leading, size),
                spaceBefore=before.pt if before is not None else 0,
                spaceAfter=after.pt if after is not None else 3,
                leftIndent=indent, bulletIndent=max(0, indent - 10), bulletFontName=fonts[0],
                bulletFontSize=size, alignment={0: 0, 1: 1, 2: 2, 3: 4}.get(setting("alignment"), 0),
                keepWithNext=is_heading or bool(setting("keep_with_next", False)),
            )
        story.append(Paragraph(text, style, bulletText="•" if bullet else None))
        previous_kind = kind
        borders = paragraph._p.xpath("./w:pPr/w:pBdr/w:bottom")
        if resume and is_heading:
            rule = HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#D8E2DC"), spaceAfter=7)
            rule.keepWithNext = True
            story.append(rule)
        elif not resume and borders and borders[0].get(qn("w:val")) not in {"nil", "none"}:
            border = borders[0]
            color = border.get(qn("w:color"), "000000")
            rule = HRFlowable(width="100%", thickness=int(border.get(qn("w:sz"), "6")) / 8,
                              color=colors.HexColor(f"#{color if color != 'auto' else '000000'}"), spaceAfter=3)
            rule.keepWithNext = is_heading
            story.append(rule)
    pdf.build(story)
