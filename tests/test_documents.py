from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader

from jobbing.documents import save_pdf, save_resume
from tests.test_generate_application import draft_response


def hyperlink(paragraph, text, url):
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = text
    run.append(value)
    element.append(run)
    paragraph._p.append(element)


class DocumentTests(unittest.TestCase):
    def test_original_header_green_accents_margins_and_links_survive_pdf(self):
        source = Document()
        # The real source DOCX has no default character style; run.style can be None.
        default_character = source.styles["Default Paragraph Font"]._element
        default_character.getparent().remove(default_character)
        source.sections[0].left_margin = Inches(0.7)
        source.styles["Normal"].font.name = "Calibri"
        source.styles["Normal"].font.size = Pt(10.5)
        source.styles["Heading 1"].font.color.rgb = RGBColor.from_string("1E8E3E")
        source.styles["Heading 1"].font.size = Pt(12)
        name = source.add_paragraph("Original Candidate").runs[0]
        name.font.size = Pt(22)
        name.bold = True
        subtitle = source.add_paragraph("Software Engineer").runs[0]
        subtitle.font.color.rgb = RGBColor.from_string("1E8E3E")
        header = source.add_paragraph()
        hyperlink(header, "Profile", "https://example.com/profile?a=1&b=2")
        heading = source.add_heading("Summary", level=1)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "6"), ("color", "1E8E3E")):
            bottom.set(qn(f"w:{key}"), value)
        border.append(bottom)
        heading._p.get_or_add_pPr().append(border)
        source.add_paragraph("Original summary")
        source.add_heading("Projects", level=1)
        hyperlink(source.add_paragraph(), "Project", "https://example.com/project")
        hyperlink(source.add_paragraph(), "Other project", "https://example.com/other")
        original = BytesIO()
        source.save(original)
        draft = draft_response().resume
        draft.sections[0].paragraphs.append("Project: Python & SQL <APIs>")
        with TemporaryDirectory() as folder:
            output, pdf = Path(folder) / "resume.docx", Path(folder) / "resume.pdf"
            save_resume(draft, output, original.getvalue())
            tailored = Document(output)
            self.assertEqual(tailored.paragraphs[0].text, "Original Candidate")
            self.assertEqual(tailored.paragraphs[0].runs[0].font.size.pt, 22)
            self.assertEqual(tailored.sections[0].left_margin, source.sections[0].left_margin)
            self.assertEqual(str(tailored.styles["Heading 1"].font.color.rgb), "1E8E3E")
            self.assertTrue(tailored.element.xpath('//w:pBdr/w:bottom[@w:color="1E8E3E"]'))
            links = [h for p in tailored.paragraphs for h in p.hyperlinks]
            self.assertEqual({h.text for h in links}, {"Profile", "Project", "Other project"})
            self.assertEqual(len(links), 3)  # Inline restored link is not duplicated in Links section.
            save_pdf(output, pdf)
            reader = PdfReader(pdf)
            text = "\n".join(page.extract_text() for page in reader.pages)
            self.assertIn("Original Candidate", text)
            self.assertIn("Python & SQL <APIs>", text)
            self.assertIn("Other project", text)
            targets = {annotation.get_object()["/A"]["/URI"] for page in reader.pages
                       for annotation in page.get("/Annots", []) if "/A" in annotation.get_object()}
            self.assertEqual(targets, {"https://example.com/profile?a=1&b=2", "https://example.com/project",
                                       "https://example.com/other"})
            green = [30 / 255, 142 / 255, 62 / 255]
            operators = [args for page in reader.pages for args, op in page.get_contents().operations
                         if op in (b"rg", b"RG")]
            self.assertTrue(any(all(abs(float(a) - b) < 0.00001 for a, b in zip(args, green))
                                for args in operators))

    def test_experience_headers_are_bold_in_docx_and_pdf(self):
        source = Document()
        source.add_paragraph("Original Candidate")
        source.add_heading("Summary", 1)
        source.add_paragraph("Python engineer")
        # Exercise both exact-source cloning and newly worded role header lines.
        unchanged = "ACE RegTech (RegAI), Amsterdam, Netherlands"
        source.add_paragraph(unchanged).runs[0].bold = False
        original = BytesIO()
        source.save(original)
        draft = draft_response().resume
        draft.sections[0].heading = "Professional Experience"
        headers = ["Senior Full Stack Engineer, ACE RegTech (RegAI), Amsterdam, Netherlands, JAN 2026 – PRESENT",
                   unchanged, "JAN 2026 – PRESENT"]
        draft.sections[0].paragraphs = [*headers, "- Built Python APIs with SQL."]
        with TemporaryDirectory() as folder:
            output, pdf = Path(folder) / "resume.docx", Path(folder) / "resume.pdf"
            save_resume(draft, output, original.getvalue())
            document = Document(output)
            for text in headers:
                paragraph = next(p for p in document.paragraphs if p.text == text.replace("–", "-"))
                self.assertTrue(all(r.bold for r in paragraph.runs if r.text))
            bullet = next(p for p in document.paragraphs if p.text == "Built Python APIs with SQL.")
            self.assertFalse(any(r.bold for r in bullet.runs))
            save_pdf(output, pdf)
            rendered = []
            for page in PdfReader(pdf).pages:
                page.extract_text(visitor_text=lambda text, cm, tm, font, size: rendered.append(
                    (text, str(font.get("/BaseFont", "")) if font else "")))
            for phrase in ("Senior Full Stack Engineer", "ACE RegTech", "JAN 2026"):
                matching = [font for text, font in rendered if phrase in text]
                self.assertTrue(matching, phrase)
                self.assertTrue(all("bold" in font.lower() for font in matching), matching)
            matching = [font for text, font in rendered if "Built Python APIs" in text]
            self.assertTrue(matching)
            self.assertTrue(all("bold" not in font.lower() for font in matching))

    def test_plain_punctuation_in_both_exports_preserves_links_and_formatting(self):
        source = Document()
        source.add_paragraph("José O’Neil")
        url = "https://example.com/portfolio?q=2020–2024&label=“AI”"
        hyperlink(source.add_paragraph(), "Portfolio—AI", url)
        source.add_heading("Summary", 1)
        source.add_paragraph("Original summary")
        unchanged = "Built “AI” tools — shipped in 2020–2024…"
        source.add_paragraph(unchanged).runs[0].bold = True
        original = BytesIO()
        source.save(original)
        draft = draft_response().resume
        draft.summary = "Python engineer — builds AI/LLM tools…"
        draft.sections[0].heading = "Projects"
        draft.sections[0].paragraphs = [unchanged, "- Improved Python–SQL APIs."]
        with TemporaryDirectory() as folder:
            docx, pdf = Path(folder) / "resume.docx", Path(folder) / "resume.pdf"
            save_resume(draft, docx, original.getvalue())
            tailored = Document(docx)
            text = "\n".join(p.text for p in tailored.paragraphs)
            self.assertNotRegex(text, "[—–‘’“”…]")
            self.assertIn("José O'Neil", text)
            self.assertIn("Python engineer - builds AI/LLM tools...", text)
            self.assertIn("Improved Python-SQL APIs.", text)
            cloned = next(p for p in tailored.paragraphs if p.text == 'Built "AI" tools - shipped in 2020-2024...')
            self.assertTrue(cloned.runs[0].bold)
            self.assertEqual(tailored.paragraphs[1].hyperlinks[0].text, "Portfolio-AI")
            self.assertEqual(tailored.paragraphs[1].hyperlinks[0].url, url)
            # Also cover rendering an older DOCX directly, without rewriting it on disk.
            old_docx = Path(folder) / "original.docx"
            old_docx.write_bytes(original.getvalue())
            for path in (docx, old_docx):
                before = path.read_bytes()
                save_pdf(path, pdf)
                self.assertEqual(path.read_bytes(), before)
                reader = PdfReader(pdf)
                text = "\n".join(p.extract_text() for p in reader.pages)
                self.assertNotRegex(text, "[—–‘’“”…]")
                self.assertIn("José O'Neil", text)
                self.assertIn('Built "AI" tools - shipped in 2020-2024...', text)
                targets = {ref.get_object()["/A"]["/URI"] for p in reader.pages
                           for ref in p.get("/Annots", [])}
                self.assertEqual(targets, {url})

    def test_resume_typography_spacing_and_embedded_fonts(self):
        source = Document()
        source.styles["Normal"].font.size = Pt(7)
        source.styles["Normal"].paragraph_format.line_spacing = 1
        source.add_paragraph("Alex Morgan")
        source.add_paragraph("Software Engineer").runs[0].font.size = Pt(13)
        contact = source.add_paragraph("Amsterdam | ")
        hyperlink(contact, "alex@example.com", "mailto:alex@example.com")
        hyperlink(source.add_paragraph(), "https://example.com/" + "long-profile-" * 18,
                  "https://example.com/profile")
        source.add_heading("Professional Summary", 1)
        source.add_paragraph("Readable summary with Python & SQL <APIs>.")
        source.add_heading("Technical Skills", 1)
        skills = source.add_paragraph()
        for text in ("Back", "end:", " Python, SQL"):
            skills.add_run(text)
        hyperlink(source.add_paragraph(), "Unsafe link", "javascript:alert(1)")
        with TemporaryDirectory() as folder:
            docx, output = Path(folder) / "resume.docx", Path(folder) / "resume.pdf"
            source.save(docx)
            original = docx.read_bytes()
            save_pdf(docx, output)
            self.assertEqual(docx.read_bytes(), original)
            pdf = PdfReader(output)
            self.assertEqual(len(pdf.pages), 1)
            page = pdf.pages[0]
            rendered = []
            text = page.extract_text(visitor_text=lambda text, cm, tm, font, size: rendered.append(
                (text, size, str(font.get("/BaseFont", "")) if font else "", cm[4] + tm[4])))
            for phrase, size in (("Alex Morgan", 28), ("Software Engineer", 12),
                                 ("alex@example.com", 9.5), ("Professional Summary", 11),
                                 ("Readable summary", 10.5)):
                matches = [s for t, s, _, _ in rendered if phrase in t]
                self.assertTrue(matches, phrase)
                self.assertTrue(all(s == size for s in matches), (phrase, matches))
            self.assertTrue(all("bold" in font.lower() for t, _, font, _ in rendered
                                if "Back" in t or "end:" in t))
            self.assertTrue(any("Python, SQL" in t and "bold" not in font.lower()
                                for t, _, font, _ in rendered))
            summary_x = next(x for t, _, _, x in rendered if "Readable summary" in t)
            self.assertAlmostEqual(summary_x, 50.4)
            leading = [float(args[0]) for args, op in page.get_contents().operations if op == b"TL"]
            self.assertIn(14.25, leading)
            self.assertIn("Python & SQL <APIs>", text)
            self.assertTrue(text.strip().endswith("1"))
            self.assertLess(text.index("Alex Morgan"), text.index("Professional Summary"))
            self.assertLess(text.index("Professional Summary"), text.index("Technical Skills"))
            embedded = [font.get_object() for font in page["/Resources"]["/Font"].values()
                        if font.get_object().get("/Subtype") == "/TrueType"]
            self.assertGreaterEqual(len(embedded), 2)
            for font in embedded:
                self.assertIn("/ToUnicode", font)
                self.assertIn("/FontFile2", font["/FontDescriptor"])
            targets = set()
            for reference in page["/Annots"]:
                annotation = reference.get_object()
                targets.add(annotation["/A"]["/URI"])
                x1, _, x2, _ = map(float, annotation["/Rect"])
                self.assertGreaterEqual(x1, 50.3)
                self.assertLessEqual(x2, float(page.mediabox.width) - 50.3)
            self.assertEqual(targets, {"mailto:alex@example.com", "https://example.com/profile"})

    def test_section_and_role_headers_stay_with_first_achievement(self):
        document = Document()
        document.sections[0].page_height = Inches(5)
        for i in range(12):
            document.add_heading(f"Experience [{i:02d}]", 1)
            document.add_paragraph(f"Engineer [{i:02d}]")
            document.add_paragraph(f"Employer [{i:02d}] | 2020–Present")
            document.add_paragraph(f"Achievement [{i:02d}]: Built reliable APIs.", "List Bullet")
        with TemporaryDirectory() as folder:
            source, output = Path(folder) / "resume.docx", Path(folder) / "resume.pdf"
            document.save(source)
            save_pdf(source, output)
            pages = [page.extract_text() for page in PdfReader(output).pages]
            self.assertGreater(len(pages), 1)
            for i in range(12):
                containing = [text for text in pages if f"Experience [{i:02d}]" in text]
                self.assertEqual(len(containing), 1)
                for phrase in ("Engineer", "Employer", "Achievement"):
                    self.assertIn(f"{phrase} [{i:02d}]", containing[0])

    def test_multipage_pdf_keeps_extractable_text(self):
        document = Document()
        document.add_heading("Experience", 1)
        for i in range(90):
            document.add_paragraph(f"Role {i}: Built Python APIs and SQL services. " * 5)
        with TemporaryDirectory() as folder:
            source, output = Path(folder) / "long.docx", Path(folder) / "long.pdf"
            document.save(source)
            save_pdf(source, output)
            pdf = PdfReader(output)
            self.assertGreater(len(pdf.pages), 1)
            self.assertIn("Role 89", "\n".join(page.extract_text() for page in pdf.pages))


if __name__ == "__main__":
    unittest.main()
