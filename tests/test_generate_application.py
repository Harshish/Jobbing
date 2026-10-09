import csv
from hashlib import sha256
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

import jobbing
from jobbing import generate_application as generator
from tests.test_workflow import APP, sample_jobs, write_resume


def draft_response(with_letter=True):
    return generator.ApplicationDraft(
        resume=generator.ResumeDraft(
            name="Test Candidate", contact_lines=["candidate@example.com"],
            summary="Python engineer with SQL experience.",
            sections=[generator.ResumeSection(heading="Experience", paragraphs=[
                "Engineer | Example Employer | 2020–Present", "- Built Python APIs with SQL.",
            ])],
        ),
        motivation_letter=["Dear Hiring Team,", "My Python and SQL experience aligns with this role.",
                           "Sincerely, Test Candidate"] if with_letter else None,
        supported_keywords=["Python", "SQL"], missing_requirements=["Cloud experience not evidenced"],
        review_notes=["Check every rewritten bullet before submission."],
    )


def fake_response(draft=None, refusal=None):
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(parsed=draft, refusal=refusal))])


class GenerationTests(unittest.TestCase):
    def test_docx_outputs_provenance_and_versions(self):
        with TemporaryDirectory() as folder, patch.object(generator, "OpenAI") as client, \
                patch.dict("os.environ", {"OPENAI_API_KEY": "offline-test"}):
            root = Path(folder)
            original = root / "original.docx"
            write_resume(original)
            before = original.read_bytes()
            job = sample_jobs().to_dict("records")[0]
            job["llm_reason"] = "Unverified prior assessment"
            parse = client.return_value.__enter__.return_value.chat.completions.parse
            parse.return_value = fake_response(draft_response())
            output = generator.generate_application(original, job, root / "applications", include_letter=True)
            self.assertEqual(original.read_bytes(), before)
            self.assertTrue((output / "motivation_letter.docx").exists())
            self.assertIn("Test Candidate", PdfReader(output / "resume.pdf").pages[0].extract_text())
            letter_page = PdfReader(output / "motivation_letter.pdf").pages[0]
            letter_sizes = []
            letter_text = letter_page.extract_text(visitor_text=lambda text, cm, tm, font, size:
                                                  letter_sizes.append(size) if text.strip() else None)
            self.assertIn("Dear Hiring Team", letter_text)
            self.assertTrue(letter_text.strip().endswith("Sincerely, Test Candidate"))
            self.assertEqual(set(letter_sizes), {11})
            document = Document(output / "resume.docx")
            self.assertFalse(document.tables)
            self.assertEqual(document.paragraphs[0].text, "Test Candidate")
            self.assertTrue(any(p.style.name == "List Bullet" for p in document.paragraphs))
            self.assertEqual(document.styles["Normal"].font.name, "Arial")
            self.assertIn("Dear Hiring Team,", generator.read_resume(output / "motivation_letter.docx"))
            self.assertIn("No ATS score", (output / "review_notes.txt").read_text())
            metadata = json.loads((output / "application.json").read_text())
            self.assertEqual(metadata["resume_sha256"], sha256(before).hexdigest())
            self.assertNotIn("llm_reason", metadata["job"])
            payload = json.loads(parse.call_args.kwargs["messages"][1]["content"])
            self.assertIn("Python engineer with SQL experience", payload["original_resume"])
            self.assertEqual(payload["job"]["title"], job["title"])
            self.assertNotIn("llm_reason", payload["job"])
            self.assertTrue(payload["include_letter"])

            # Even if the model returns an unrequested letter, do not save it.
            second = generator.generate_application(before, job, root / "applications", include_letter=False)
            self.assertNotEqual(output, second)
            self.assertEqual(output.parent, second.parent)
            self.assertFalse((second / "motivation_letter.docx").exists())
            self.assertFalse((second / "motivation_letter.pdf").exists())
            self.assertTrue((second / "resume.pdf").exists())
            self.assertTrue((output / "motivation_letter.docx").exists())
            self.assertEqual(json.loads((second / "application.json").read_text())["resume_source"], "uploaded DOCX")

    def test_refusal_missing_letter_and_render_failure_publish_nothing(self):
        with TemporaryDirectory() as folder, patch.object(generator, "OpenAI") as client, \
                patch.dict("os.environ", {"OPENAI_API_KEY": "offline-test"}):
            root = Path(folder)
            original = root / "original.docx"
            write_resume(original)
            job = sample_jobs().to_dict("records")[0]
            parse = client.return_value.__enter__.return_value.chat.completions.parse
            for response in (fake_response(refusal="Cannot comply"), fake_response(draft_response(False))):
                parse.return_value = response
                with self.assertRaises(ValueError):
                    generator.generate_application(original, job, root / "applications", include_letter=True)
            parse.return_value = fake_response(draft_response())
            for renderer in ("save_resume", "save_pdf"):
                with patch.object(generator, renderer, side_effect=OSError("Disk full")):
                    with self.assertRaisesRegex(OSError, "Disk full"):
                        generator.generate_application(original, job, root / "applications")
            self.assertFalse(list((root / "applications").rglob("resume.docx")))
            self.assertFalse(list((root / "applications").rglob(".draft-*")))

    def test_job_validation_and_safe_grouping(self):
        with patch.object(generator, "OpenAI") as client:
            with self.assertRaisesRegex(ValueError, "description"):
                generator.generate_application("not-needed.docx", {"title": "Engineer", "description": ""})
            client.assert_not_called()
        job = {"company": "../../outside/", "title": "../role", "job_url": "https://example.com/1"}
        name = generator.job_folder_name(job)
        self.assertNotIn("/", name)
        self.assertNotIn("..", name)
        self.assertEqual(name, generator.job_folder_name({**job, "llm_match_score": "95"}))
        self.assertNotEqual(name, generator.job_folder_name({**job, "job_url": "https://example.com/2"}))

    def test_cli_selects_the_requested_matched_row(self):
        with TemporaryDirectory() as folder, patch.object(generator, "generate_application") as generate:
            root = Path(folder)
            matches = root / "matched_jobs.csv"
            rows = sample_jobs()
            rows["llm_match_score"] = [85, 95]
            rows.to_csv(matches, index=False, escapechar="\\", quoting=csv.QUOTE_ALL)
            generate.return_value = root / "applications" / "draft"
            argv = ["generate_application", "--matches", str(matches), "--job-number", "2", "--with-letter"]
            with patch("sys.argv", argv):
                generator.main()
            self.assertEqual(generate.call_args.args[1]["title"], "Platform Engineer")
            self.assertTrue(generate.call_args.kwargs["include_letter"])
            with patch("sys.argv", [*argv[:], "--job-number", "3"]), self.assertRaises(SystemExit):
                generator.main()
            self.assertEqual(generate.call_count, 1)


class ApplicationWebTests(unittest.TestCase):
    def test_buttons_downloads_persistence_and_no_generation_on_browse(self):
        with TemporaryDirectory() as folder, patch.object(jobbing, "DATA", Path(folder)), \
                patch.object(generator, "OpenAI") as client, \
                patch.dict("os.environ", {"OPENAI_API_KEY": "offline-test"}):
            root = Path(folder)
            write_resume(root / "resume.docx")
            matches = root / "matched_jobs.csv"
            rows = sample_jobs()
            rows["llm_match_score"] = [85, 95]
            rows.to_csv(matches, index=False, escapechar="\\", quoting=csv.QUOTE_ALL)
            parse = client.return_value.__enter__.return_value.chat.completions.parse
            parse.return_value = fake_response(draft_response())
            at = AppTest.from_file(str(APP), default_timeout=20).run()
            key = f"generate:{matches}:1"
            self.assertFalse(at.exception)
            self.assertFalse(at.button(key=key).disabled)
            client.assert_not_called()
            at.checkbox(key="include_letter").check().run()
            at.button(key=key).click().run()
            self.assertFalse(at.exception)
            self.assertFalse(at.error)
            self.assertEqual(client.call_count, 1)
            self.assertEqual(len(list(root.glob("applications/*/*/resume.docx"))), 1)
            self.assertEqual(len(list(root.glob("applications/*/*/motivation_letter.docx"))), 1)
            self.assertEqual(len(at.get("download_button")), 5)  # CSV, two PDFs, two editable DOCX files.
            at.text_input(key="search_results").set_value("Example One").run()
            self.assertFalse(at.exception)
            self.assertEqual(client.call_count, 1)
            self.assertEqual(len(at.get("download_button")), 5)

            # Saved downloads survive a fresh browser session; no API call is made.
            fresh = AppTest.from_file(str(APP), default_timeout=20).run()
            self.assertFalse(fresh.exception)
            self.assertEqual(len(fresh.get("download_button")), 5)
            self.assertEqual(client.call_count, 1)

            # Failure shows an error and preserves previous successful documents.
            parse.side_effect = RuntimeError("Provider offline")
            at.button(key=key).click().run()
            self.assertFalse(at.exception)
            self.assertTrue(any("Provider offline" in error.value for error in at.error))
            self.assertEqual(len(list(root.glob("applications/*/*/resume.docx"))), 1)
            self.assertEqual(len(at.get("download_button")), 5)


if __name__ == "__main__":
    unittest.main()
